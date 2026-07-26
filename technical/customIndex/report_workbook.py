"""Create the dated custom-index breakout workbook after the ES enrichment run."""

from __future__ import annotations

import argparse
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterable
from zoneinfo import ZoneInfo

import pandas as pd
from elasticsearch import Elasticsearch
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

try:  # Package import for tests and module execution.
    from .custom_index_pipeline import DEFAULT_DEFINITION_FILE, load_index_definitions, normalise_ticker
    from .fundamentals import (
        DEFAULT_REQUEST_DELAY_SECONDS,
        FundamentalMetrics,
        fetch_constituent_fundamentals,
    )
    from .report_breakouts import ES_HOST, find_breakouts, latest_custom_index_date
except ImportError:  # Direct script execution from run.sh.
    from custom_index_pipeline import DEFAULT_DEFINITION_FILE, load_index_definitions, normalise_ticker
    from fundamentals import (
        DEFAULT_REQUEST_DELAY_SECONDS,
        FundamentalMetrics,
        fetch_constituent_fundamentals,
    )
    from report_breakouts import ES_HOST, find_breakouts, latest_custom_index_date


IST = ZoneInfo("Asia/Kolkata")
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parents[2] / "outputs" / "custom-index-breakouts"
TAG_COLUMNS = (
    "Screener Sub-Sector",
    "Primary Sub-Sector",
    "Secondary Tags",
    "Broad Market Indices (NIFTY ES tags)",
)
NO_TAG_PREFIXES = ("no local ", "no stored ", "no matching ")

TITLE_FILL = "0B1F3A"
HEADER_FILL = "1F4E78"
CARD_FILL = "EAF3F5"
ALT_FILL = "D9EAF7"
LIGHT_BORDER = "D9E2E8"
WHITE = "FFFFFF"


@dataclass
class ConstituentRow:
    """One de-duplicated stock present in at least one breakout basket."""

    ticker: str
    tradingview_symbol: str
    company_name: str = ""
    breakout_indices: list[str] = field(default_factory=list)
    tags: dict[str, list[str]] = field(
        default_factory=lambda: {column: [] for column in TAG_COLUMNS}
    )


def report_as_of_date(now: datetime | None = None) -> date:
    """Use the live weekday, or the preceding Friday after the week has closed."""
    local_now = now or datetime.now(IST)
    if local_now.tzinfo is None:
        local_now = local_now.replace(tzinfo=IST)
    else:
        local_now = local_now.astimezone(IST)

    local_date = local_now.date()
    if local_date.weekday() <= 4:
        return local_date
    return local_date - timedelta(days=local_date.weekday() - 4)


def _display_ticker(normalised_ticker: str) -> str:
    return normalised_ticker.removesuffix(".NS").removesuffix(".BO")


def tradingview_symbol(normalised_ticker: str) -> str:
    """Translate the Yahoo symbol stored in the index definition for TradingView."""
    if normalised_ticker.endswith(".BO"):
        return f"BSE:{normalised_ticker.removesuffix('.BO')}"
    return f"NSE:{normalised_ticker.removesuffix('.NS')}"


def tradingview_formula(constituents: Iterable[str]) -> str:
    """Return a copy/pasteable equal-price TradingView basket expression."""
    symbols = [tradingview_symbol(ticker) for ticker in constituents]
    return f"({' + '.join(symbols)})/{len(symbols)}" if symbols else ""


def _normalise_tag(value: object) -> str:
    if value is None or pd.isna(value):
        return ""
    tag = str(value).strip()
    if not tag or tag.casefold().startswith(NO_TAG_PREFIXES):
        return ""
    return tag


def collect_constituents(
    breakout_documents: Iterable[dict], definition_file: Path = DEFAULT_DEFINITION_FILE
) -> list[ConstituentRow]:
    """Read checked-in constituent tags for every breakout basket."""
    by_ticker: OrderedDict[str, ConstituentRow] = OrderedDict()
    definitions = {
        definition.definition_id: definition
        for definition in load_index_definitions(definition_file)
    }

    for document in breakout_documents:
        index_name = document.get("index_name") or document.get("ticker", "")
        definition_id = str(document.get("definition_id", ""))
        definition = definitions.get(definition_id)
        if definition is None:
            print(f"Skipping constituent drill-down for {index_name}: definition is missing")
            continue

        for detail in definition.constituent_details:
            normalised_ticker = normalise_ticker(detail.get("yahoo_symbol"))
            if normalised_ticker is None:
                continue

            constituent = by_ticker.setdefault(
                normalised_ticker,
                ConstituentRow(
                    ticker=_display_ticker(normalised_ticker),
                    tradingview_symbol=tradingview_symbol(normalised_ticker),
                ),
            )
            if index_name not in constituent.breakout_indices:
                constituent.breakout_indices.append(index_name)

            if not constituent.company_name:
                constituent.company_name = _normalise_tag(detail.get("company_name"))

            details_tags = detail.get("tags", {})
            for tag_column in TAG_COLUMNS:
                tag = _normalise_tag(details_tags.get(tag_column))
                if tag and tag not in constituent.tags[tag_column]:
                    constituent.tags[tag_column].append(tag)

    return list(by_ticker.values())


def nearest_level(document: dict) -> dict:
    """Return the ES support/resistance level that qualified the screen."""
    return next(
        (
            level
            for level in document.get("crossed_resistance", [])
            if level.get("support_distance_pct") is not None
            and level["support_distance_pct"] <= 10
        ),
        {},
    )


def _to_excel_pct(value: object) -> float | None:
    return float(value) / 100 if value is not None else None


def breakout_table_rows(documents: Iterable[dict]) -> list[list[object]]:
    rows: list[list[object]] = []
    for document in documents:
        level = nearest_level(document)
        close = document.get("close")
        resistance = level.get("resistance_level")
        if resistance is None:
            resistance_state = "No overhead level"
        elif close is not None and close > resistance:
            resistance_state = "Confirmed breakout"
        else:
            resistance_state = "Below resistance"

        rows.append(
            [
                document.get("index_name", document.get("ticker", "")),
                document.get("tradingview_equation") or tradingview_formula(document.get("constituents", [])),
                "",  # Filled after source-workbook tags are collected below.
                "",
                "",
                "",
                document.get("date"),
                close,
                level.get("support_level"),
                _to_excel_pct(level.get("support_distance_pct")),
                resistance,
                _to_excel_pct(level.get("resistance_distance_pct")),
                _to_excel_pct(document.get("dist_from_52w_high_pct")),
                document.get("rsi"),
                _to_excel_pct(document.get("roc")),
                document.get("constituent_count", len(document.get("constituents", []))),
                "equal_weight (daily rebalanced)",
                resistance_state,
            ]
        )
    return rows


def _index_tag_unions(constituents: Iterable[ConstituentRow]) -> dict[str, dict[str, str]]:
    index_tags: dict[str, dict[str, list[str]]] = {}
    for constituent in constituents:
        for index_name in constituent.breakout_indices:
            bucket = index_tags.setdefault(
                index_name, {tag_column: [] for tag_column in TAG_COLUMNS}
            )
            for tag_column, tags in constituent.tags.items():
                for tag in tags:
                    for tag_part in (part.strip() for part in tag.split("|")):
                        if tag_part and tag_part not in bucket[tag_column]:
                            bucket[tag_column].append(tag_part)

    return {
        index_name: {
            tag_column: " | ".join(tags)
            for tag_column, tags in tags_by_column.items()
        }
        for index_name, tags_by_column in index_tags.items()
    }


def _joined_tags(constituent: ConstituentRow, tag_column: str) -> str:
    return " | ".join(constituent.tags[tag_column])


def fundamentals_table_rows(
    constituents: Iterable[ConstituentRow],
    fundamentals_by_ticker: dict[str, FundamentalMetrics],
) -> list[list[object]]:
    """Combine report-only Screener data with the checked-in constituent taxonomy."""
    rows: list[list[object]] = []
    for constituent in constituents:
        metrics = fundamentals_by_ticker.get(
            constituent.ticker,
            FundamentalMetrics(ticker=constituent.ticker, fetch_status="Unavailable: no response"),
        )
        rows.append(
            [
                constituent.ticker,
                constituent.tradingview_symbol,
                constituent.company_name,
                " | ".join(constituent.breakout_indices),
                metrics.market_cap_cr,
                metrics.stock_pe,
                _to_excel_pct(metrics.roce_pct),
                _to_excel_pct(metrics.roe_pct),
                _to_excel_pct(metrics.revenue_qoq_pct),
                _to_excel_pct(metrics.revenue_yoy_pct),
                _to_excel_pct(metrics.profit_qoq_pct),
                _to_excel_pct(metrics.profit_yoy_pct),
                "Yes" if metrics.current_quarter_result_out else "No"
                if metrics.current_quarter_result_out is not None
                else "",
                metrics.expected_result_quarter,
                metrics.latest_reported_quarter,
                _joined_tags(constituent, "Broad Market Indices (NIFTY ES tags)"),
                " | ".join(constituent.breakout_indices),
                _joined_tags(constituent, "Screener Sub-Sector"),
                metrics.screener_sector,
                metrics.screener_industry,
                _joined_tags(constituent, "Primary Sub-Sector"),
                _joined_tags(constituent, "Secondary Tags"),
                metrics.screener_url,
                metrics.fetch_status,
            ]
        )
    return rows


def _apply_title(sheet, title: str, subtitle: str, end_column: str) -> None:
    sheet.merge_cells(f"A1:{end_column}1")
    title_cell = sheet["A1"]
    title_cell.value = title
    title_cell.font = Font(bold=True, color=WHITE, size=16)
    title_cell.fill = PatternFill("solid", fgColor=TITLE_FILL)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(f"A2:{end_column}2")
    subtitle_cell = sheet["A2"]
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(italic=True, color="52616B")
    subtitle_cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[2].height = 20
    sheet.sheet_view.showGridLines = False


def _apply_header(sheet, row: int, start_column: str, end_column: str) -> None:
    thin = Side(style="thin", color=WHITE)
    for cell in sheet[f"{start_column}{row}:{end_column}{row}"][0]:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    sheet.row_dimensions[row].height = 34


def _apply_data_layout(sheet, first_row: int, last_row: int, columns: int) -> None:
    if last_row < first_row:
        return
    thin = Side(style="thin", color=LIGHT_BORDER)
    for row in sheet.iter_rows(
        min_row=first_row, max_row=last_row, min_col=1, max_col=columns
    ):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        if row[0].row % 2:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=ALT_FILL)
        sheet.row_dimensions[row[0].row].height = 45


def _add_table(sheet, name: str, start_row: int, end_row: int, end_column: str) -> None:
    if end_row <= start_row:
        return
    table = Table(displayName=name, ref=f"A{start_row}:{end_column}{end_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    sheet.add_table(table)


def build_workbook(
    documents: list[dict],
    constituents: list[ConstituentRow],
    candle_date: str,
    as_of: date,
    fundamentals_by_ticker: dict[str, FundamentalMetrics] | None = None,
) -> Workbook:
    """Build a self-contained Excel report from the latest ES screen."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    breakouts_sheet = workbook.create_sheet("Breakouts")
    constituents_sheet = workbook.create_sheet("Breakout Constituents")
    methodology_sheet = workbook.create_sheet("Methodology")
    fundamentals_sheet = workbook.create_sheet("Fundamentals")

    index_tags = _index_tag_unions(constituents)
    breakout_rows = breakout_table_rows(documents)
    for row in breakout_rows:
        row[2:6] = [index_tags.get(row[0], {}).get(tag_column, "") for tag_column in TAG_COLUMNS]

    _apply_title(
        breakouts_sheet,
        "Custom Index Breakout Monitor",
        f"Report as of {as_of.isoformat()} • Weekly candle beginning {candle_date} • Equal-weight, daily-rebalanced synthetic indices",
        "R",
    )
    breakouts_sheet["A3"] = "Signal count"
    breakouts_sheet["B3"] = len(breakout_rows)
    breakouts_sheet["D3"] = "Sector signals"
    breakouts_sheet["E3"] = sum(1 for document in documents if document.get("classification_type") == "sector")
    breakouts_sheet["G3"] = "Industry signals"
    breakouts_sheet["H3"] = sum(1 for document in documents if document.get("classification_type") == "industry")
    breakouts_sheet["J3"] = "Candle week"
    breakouts_sheet["K3"] = datetime.strptime(candle_date, "%Y-%m-%d").date()
    breakouts_sheet["K3"].number_format = "yyyy-mm-dd"
    for cell in ("A3", "D3", "G3", "J3"):
        breakouts_sheet[cell].fill = PatternFill("solid", fgColor=CARD_FILL)
        breakouts_sheet[cell].font = Font(bold=True, color="16324F")
    for cell in ("B3", "E3", "H3", "K3"):
        breakouts_sheet[cell].fill = PatternFill("solid", fgColor=WHITE)
        breakouts_sheet[cell].font = Font(bold=True)

    breakouts_sheet.merge_cells("A4:R4")
    breakouts_sheet["A4"] = (
        "Screen: VCP trend template is true and the nearest support is within 10% "
        "of the latest weekly close. Resistance State indicates the position versus "
        "the nearest resistance; tag fields are the de-duplicated union of constituent tags."
    )
    breakouts_sheet["A4"].fill = PatternFill("solid", fgColor="F3F6F8")
    breakouts_sheet["A4"].font = Font(italic=True, color="43515C")
    breakouts_sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    breakouts_sheet.row_dimensions[4].height = 30

    breakout_headers = [
        "Index Name", "TradingView Formula", "Screener Sub-Sector", "Primary Sub-Sector",
        "Secondary Tags", "Broad Market Indices (NIFTY ES tags)", "Week Start", "Close",
        "Support", "Support Gap %", "Resistance", "Resistance Gap %", "52W High Gap %",
        "RSI (14)", "ROC (20)", "Constituents", "Calculation", "Resistance State",
    ]
    breakouts_sheet.append([])
    breakouts_sheet.append(breakout_headers)
    for row in breakout_rows:
        breakouts_sheet.append(row)
    _apply_header(breakouts_sheet, 6, "A", "R")
    _apply_data_layout(breakouts_sheet, 7, 6 + len(breakout_rows), len(breakout_headers))
    for row in range(7, 7 + len(breakout_rows)):
        breakouts_sheet.row_dimensions[row].height = 72
    _add_table(breakouts_sheet, "CustomIndexBreakouts", 6, 6 + len(breakout_rows), "R")
    if not breakout_rows:
        breakouts_sheet.merge_cells("A7:R7")
        breakouts_sheet["A7"] = f"No custom indices matched the breakout screen for the week beginning {candle_date}."
        breakouts_sheet["A7"].fill = PatternFill("solid", fgColor="FFF2CC")
        breakouts_sheet["A7"].font = Font(italic=True, color="9C6500")
        breakouts_sheet["A7"].alignment = Alignment(vertical="center")
        breakouts_sheet.row_dimensions[7].height = 24
    breakouts_sheet.freeze_panes = "A7"
    for column, width in {
        "A": 24, "B": 72, "C": 38, "D": 52, "E": 48, "F": 54, "G": 14,
        "H": 14, "I": 14, "J": 15, "K": 14, "L": 17, "M": 15, "N": 10,
        "O": 12, "P": 13, "Q": 25, "R": 21,
    }.items():
        breakouts_sheet.column_dimensions[column].width = width
    for row in range(7, 7 + len(breakout_rows)):
        breakouts_sheet[f"G{row}"].number_format = "yyyy-mm-dd"
        for column in ("H", "I", "K"):
            breakouts_sheet[f"{column}{row}"].number_format = "#,##0.0;[Red](#,##0.0);-"
        for column in ("J", "L", "M", "O"):
            breakouts_sheet[f"{column}{row}"].number_format = "0.0%;[Red](0.0%);-"
        breakouts_sheet[f"N{row}"].number_format = "0.0"
    if breakout_rows:
        breakouts_sheet.conditional_formatting.add(
            f"R7:R{6 + len(breakout_rows)}",
            CellIsRule(operator="equal", formula=['"Confirmed breakout"'], fill=PatternFill("solid", fgColor="C6EFCE")),
        )

    _apply_title(
        constituents_sheet,
        "Breakout Index Constituents",
        f"Unique stocks from the {len(documents)} breakout indices • Report as of {as_of.isoformat()}",
        "H",
    )
    constituents_sheet["A3"] = "Unique stocks"
    constituents_sheet["B3"] = len(constituents)
    constituents_sheet["D3"] = "Breakout indices"
    constituents_sheet["E3"] = len(documents)
    for cell in ("A3", "D3"):
        constituents_sheet[cell].fill = PatternFill("solid", fgColor=CARD_FILL)
        constituents_sheet[cell].font = Font(bold=True, color="16324F")
    for cell in ("B3", "E3"):
        constituents_sheet[cell].fill = PatternFill("solid", fgColor=WHITE)
        constituents_sheet[cell].font = Font(bold=True)
    constituents_sheet.merge_cells("A4:H4")
    constituents_sheet["A4"] = (
        "A stock appears once even if it is part of several breakout indices; "
        "Breakout Index / Indices lists every matching basket. Blank broad-market tags "
        "mean that the checked-in basket definition has no NIFTY ES tag for that stock."
    )
    constituents_sheet["A4"].fill = PatternFill("solid", fgColor="F3F6F8")
    constituents_sheet["A4"].font = Font(italic=True, color="43515C")
    constituents_sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    constituents_sheet.row_dimensions[4].height = 32
    constituent_headers = [
        "Breakout Index / Indices", "Ticker", "TradingView Symbol", "Company Name",
        *TAG_COLUMNS,
    ]
    constituents_sheet.append([])
    constituents_sheet.append(constituent_headers)
    for constituent in constituents:
        constituents_sheet.append(
            [
                " | ".join(constituent.breakout_indices),
                constituent.ticker,
                constituent.tradingview_symbol,
                constituent.company_name,
                *(_joined_tags(constituent, tag_column) for tag_column in TAG_COLUMNS),
            ]
        )
    _apply_header(constituents_sheet, 6, "A", "H")
    _apply_data_layout(constituents_sheet, 7, 6 + len(constituents), len(constituent_headers))
    _add_table(constituents_sheet, "BreakoutConstituents", 6, 6 + len(constituents), "H")
    if not constituents:
        constituents_sheet.merge_cells("A7:H7")
        constituents_sheet["A7"] = "No constituent rows are available because the breakout screen has no matching indices."
        constituents_sheet["A7"].fill = PatternFill("solid", fgColor="FFF2CC")
        constituents_sheet["A7"].font = Font(italic=True, color="9C6500")
        constituents_sheet["A7"].alignment = Alignment(vertical="center")
        constituents_sheet.row_dimensions[7].height = 24
    constituents_sheet.freeze_panes = "A7"
    for column, width in {
        "A": 42, "B": 16, "C": 20, "D": 34, "E": 34, "F": 48, "G": 42, "H": 54,
    }.items():
        constituents_sheet.column_dimensions[column].width = width

    _apply_title(methodology_sheet, "Methodology & Data Notes", "", "B")
    methodology_sheet.append(["Item", "Detail"])
    methodology_rows = [
        ["Elasticsearch source", "nifty_custom_index_breakout on http://localhost:9200"],
        ["Report date", f"{as_of.isoformat()} — live weekday during an in-progress week; preceding Friday after the weekly close."],
        ["Candle week", f"{candle_date} (weekly candle start date stored in Elasticsearch)."],
        ["Screen", "VCP trend template = true and support distance <= 10%. This is a screened setup, not necessarily a confirmed resistance cross."],
        ["Index calculation", "Base 1000. Available constituent OHLC returns are averaged equally each day, then aggregated into weekly OHLCV. The basket is rebalanced daily."],
        ["TradingView formula", "Copy/pasteable equal-price basket expression: sum of NSE/BSE component symbols divided by constituent count. It is a visual proxy; Elasticsearch values use daily equal-weight return rebalancing."],
        ["Constituent tags", "Stock tags are stored in the checked-in TradingView basket manifest. Index-level tags are de-duplicated unions of those stock tags."],
        ["Fundamentals", "Live Screener.in values are requested only for the final breakout constituents while this workbook is created. They are not stored in Elasticsearch. QoQ compares the latest reported quarter with the prior quarter; YoY compares it with the same quarter one year earlier. Latest Quarter Result Out is Yes only when Screener has revenue, net profit, and EPS for the latest completed quarter."],
    ]
    for row in methodology_rows:
        methodology_sheet.append(row)
    _apply_header(methodology_sheet, 3, "A", "B")
    _apply_data_layout(methodology_sheet, 4, 3 + len(methodology_rows), 2)
    methodology_sheet.column_dimensions["A"].width = 24
    methodology_sheet.column_dimensions["B"].width = 110
    methodology_sheet.freeze_panes = "A4"

    fundamentals_by_ticker = fundamentals_by_ticker or {}
    fundamental_rows = fundamentals_table_rows(constituents, fundamentals_by_ticker)
    _apply_title(
        fundamentals_sheet,
        "Breakout Constituent Fundamentals",
        (
            f"Current Screener.in values for {len(constituents)} final constituents • "
            f"Report as of {as_of.isoformat()} • Not stored in Elasticsearch"
        ),
        "X",
    )
    fundamentals_sheet["A3"] = "Unique constituents"
    fundamentals_sheet["B3"] = len(constituents)
    fundamentals_sheet["D3"] = "Retrieved"
    fundamentals_sheet["E3"] = sum(
        1
        for metrics in fundamentals_by_ticker.values()
        if metrics.fetch_status.startswith("Retrieved")
    )
    fundamentals_sheet["G3"] = "Unavailable"
    fundamentals_sheet["H3"] = sum(
        1
        for metrics in fundamentals_by_ticker.values()
        if metrics.fetch_status.startswith("Unavailable")
    )
    for cell in ("A3", "D3", "G3"):
        fundamentals_sheet[cell].fill = PatternFill("solid", fgColor=CARD_FILL)
        fundamentals_sheet[cell].font = Font(bold=True, color="16324F")
    for cell in ("B3", "E3", "H3"):
        fundamentals_sheet[cell].fill = PatternFill("solid", fgColor=WHITE)
        fundamentals_sheet[cell].font = Font(bold=True)

    fundamentals_sheet.merge_cells("A4:X4")
    fundamentals_sheet["A4"] = (
        "Market Cap, Stock P/E, ROCE and ROE are current Screener values. "
        "Growth values use reported consolidated quarters. A blank metric means Screener did not "
        "provide enough usable data; use the source URL for the live company page."
    )
    fundamentals_sheet["A4"].fill = PatternFill("solid", fgColor="F3F6F8")
    fundamentals_sheet["A4"].font = Font(italic=True, color="43515C")
    fundamentals_sheet["A4"].alignment = Alignment(wrap_text=True, vertical="center")
    fundamentals_sheet.row_dimensions[4].height = 34

    fundamentals_headers = [
        "Ticker", "TradingView Symbol", "Company Name", "Breakout Index / Indices",
        "Market Cap (₹ Cr)", "Stock P/E (x)", "ROCE %", "ROE %", "Revenue QoQ %",
        "Revenue YoY %", "Profit QoQ %", "Profit YoY %", "Latest Quarter Result Out?",
        "Expected Result Quarter", "Latest Reported Quarter",
        "Broad Market Indices (NIFTY ES tags)", "Sectoral Index / Indices",
        "Screener Sub-Sector", "Screener Sector", "Screener Industry",
        "Primary Sub-Sector", "Secondary Tags", "Screener Source", "Fetch Status",
    ]
    fundamentals_sheet.append([])
    fundamentals_sheet.append(fundamentals_headers)
    for row in fundamental_rows:
        fundamentals_sheet.append(row)
    _apply_header(fundamentals_sheet, 6, "A", "X")
    _apply_data_layout(
        fundamentals_sheet, 7, 6 + len(fundamental_rows), len(fundamentals_headers)
    )
    for row in range(7, 7 + len(fundamental_rows)):
        fundamentals_sheet.row_dimensions[row].height = 48
    _add_table(
        fundamentals_sheet, "BreakoutFundamentals", 6, 6 + len(fundamental_rows), "X"
    )
    if not fundamental_rows:
        fundamentals_sheet.merge_cells("A7:X7")
        fundamentals_sheet["A7"] = "No fundamentals are available because the breakout screen has no constituent rows."
        fundamentals_sheet["A7"].fill = PatternFill("solid", fgColor="FFF2CC")
        fundamentals_sheet["A7"].font = Font(italic=True, color="9C6500")
        fundamentals_sheet["A7"].alignment = Alignment(vertical="center")
        fundamentals_sheet.row_dimensions[7].height = 24
    fundamentals_sheet.freeze_panes = "A7"
    for column, width in {
        "A": 16, "B": 20, "C": 34, "D": 42, "E": 18, "F": 14, "G": 12,
        "H": 12, "I": 15, "J": 15, "K": 15, "L": 15, "M": 17, "N": 22,
        "O": 21, "P": 52, "Q": 42, "R": 34, "S": 30, "T": 34, "U": 44,
        "V": 42, "W": 42, "X": 28,
    }.items():
        fundamentals_sheet.column_dimensions[column].width = width
    for row in range(7, 7 + len(fundamental_rows)):
        fundamentals_sheet[f"E{row}"].number_format = "#,##0.0;[Red](#,##0.0);-"
        fundamentals_sheet[f"F{row}"].number_format = "0.0x;[Red](0.0x);-"
        for column in ("G", "H", "I", "J", "K", "L"):
            fundamentals_sheet[f"{column}{row}"].number_format = "0.0%;[Red](0.0%);-"
    if fundamental_rows:
        fundamentals_sheet.conditional_formatting.add(
            f"M7:M{6 + len(fundamental_rows)}",
            CellIsRule(
                operator="equal",
                formula=['"Yes"'],
                fill=PatternFill("solid", fgColor="C6EFCE"),
            ),
        )

    return workbook


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create the custom-index breakout Excel report.")
    parser.add_argument(
        "--definition-file",
        type=Path,
        default=DEFAULT_DEFINITION_FILE,
        help="Checked-in TradingView basket definition manifest.",
    )
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument(
        "--as-of-date", type=date.fromisoformat,
        help="Override the report filename date (YYYY-MM-DD); useful for reproducible runs.",
    )
    parser.add_argument(
        "--screener-delay-seconds",
        type=float,
        default=DEFAULT_REQUEST_DELAY_SECONDS,
        help="Minimum delay between sequential Screener.in requests (default: %(default)s seconds).",
    )
    parser.add_argument(
        "--skip-fundamentals",
        action="store_true",
        help="Create the workbook without live Screener.in requests; useful for offline layout checks.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    es = Elasticsearch(ES_HOST)
    candle_date = latest_custom_index_date(es)
    if candle_date is None:
        print("No custom-index candles are indexed yet; workbook was not created.")
        return 0

    documents = find_breakouts(es, candle_date)
    constituents = collect_constituents(documents, args.definition_file)
    as_of = args.as_of_date or report_as_of_date()
    if args.skip_fundamentals:
        print("Skipping live Screener.in fundamental retrieval.")
        fundamentals_by_ticker: dict[str, FundamentalMetrics] = {}
    else:
        print(
            f"Fetching current Screener.in fundamentals for {len(constituents)} unique "
            f"breakout constituents with a {args.screener_delay_seconds:.2f}s request delay..."
        )
        fundamentals_by_ticker = fetch_constituent_fundamentals(
            (constituent.ticker for constituent in constituents),
            delay_seconds=args.screener_delay_seconds,
        )
    workbook = build_workbook(
        documents, constituents, candle_date, as_of, fundamentals_by_ticker
    )
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output_path = args.output_dir / f"custom_index_breakouts_{as_of.isoformat()}.xlsx"
    workbook.save(output_path)
    print(
        f"Created workbook: {output_path} "
        f"({len(documents)} breakout indices, {len(constituents)} unique constituents, "
        f"{sum(1 for metrics in fundamentals_by_ticker.values() if metrics.fetch_status.startswith('Retrieved'))} "
        "fundamental records retrieved)"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
